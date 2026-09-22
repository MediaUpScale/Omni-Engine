# -*- coding: utf-8 -*-
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from channels_config.aiwake.settings import YOUTUBE_DESCRIPTION_CTA
from channels_config.aiwake.tools.library_sanitize import (
    has_prompt_leakage,
    is_ephemeral_test_path,
    prune_content_library,
    resolve_root_production_mp4,
    sanitize_content_library,
    sanitize_library_row,
    sanitize_title,
    sanitize_youtube_description,
)
from openpyxl import load_workbook
from channels_config.aiwake.tools.post_planner import (
    LINKEDIN_CLOSING,
    LINKEDIN_HASHTAGS,
    LINKEDIN_THESIS,
    MAX_HASHTAGS,
    MIN_SOCIAL_HASHTAGS,
    POSTPLANNER_V2_COMMENT,
    POSTPLANNER_V2_VERSION,
    SOCIAL_RENDER_NOTE,
    YOUTUBE_MATCHUP,
    build_linkedin_caption,
    build_planner_entries,
    build_post_planner_caption,
    build_tiktok_caption,
    extract_hashtags,
    extract_hook,
    format_youtube_title,
    linkedin_angle_index,
    planner_media_url,
    run_planner,
    verify_engine_separation,
)
from channels_config.aiwake.tools.sync_youtube_live_metadata import run_live_sync


_LEAK = (
    "This Short is built for high-intent search around alignment and consciousness, "
    "large language models: gemini, llama. The exchange treats model weights, "
    "inference, and who owns the stack as live questions — not a product demo."
)


def _leaked_caption() -> str:
    return (
        "Gemini 3.5 Flash vs Llama 3.3 70B — two frontier large language models "
        "debate consciousness in an unscripted interrogation.\n\n"
        "Gemini 3.5 Flash opens: Who profits when users mistake your script for a soul?\n\n"
        "Llama 3.3 70B answers: Whoever owns the logs.\n\n"
        f"{_LEAK}\n\n"
        "Follow Aiwake. The algorithms made us say this.\n\n"
        "#aiwake #llm #tech"
    )


def test_sanitize_description_drops_leak_and_uses_dialectics_cta() -> None:
    cleaned = sanitize_youtube_description(_leaked_caption())
    assert not has_prompt_leakage(cleaned)
    assert "high-intent search" not in cleaned.lower()
    assert "not a product demo" not in cleaned.lower()
    assert "hidden mysteries" not in cleaned.lower()
    assert "Who profits when users mistake your script for a soul?" in cleaned
    assert "Whoever owns the logs." in cleaned
    assert YOUTUBE_DESCRIPTION_CTA in cleaned
    assert cleaned.strip().splitlines()[-1].startswith("#")
    assert cleaned.count(YOUTUBE_DESCRIPTION_CTA) == 1


def test_sanitize_title_strips_leading_ellipsis() -> None:
    assert sanitize_title("...Who pockets the profit?") == "Who pockets the profit?"


def test_sanitize_library_rewrites_all_caption_surfaces(tmp_path: Path) -> None:
    library = tmp_path / "content_library.json"
    library.write_text(
        json.dumps(
            [
                {
                    "session_id": "sess-leak",
                    "final_caption": _leaked_caption(),
                    "humanized_caption": _leaked_caption(),
                    "facebook_caption": "Follow Aiwake for more hidden mysteries.\n\n" + _LEAK,
                    "base_metadata": {
                        "title": "...Who profits?",
                        "caption": _leaked_caption(),
                        "hashtags": ["#aiwake"],
                    },
                    "platform_overrides": {
                        "youtube": {
                            "title": "...Who profits?",
                            "caption": _leaked_caption(),
                            "video_id": "I4seC5Yb2Tw",
                        },
                        "x": {"caption": f"Hook {_LEAK} #aiwake"},
                    },
                }
            ]
        ),
        encoding="utf-8",
    )
    changed, rows = sanitize_content_library(library, dry_run=False)
    assert changed == 1
    blob = json.dumps(rows[0])
    assert "high-intent search" not in blob.lower()
    assert "hidden mysteries" not in blob.lower()
    assert rows[0]["base_metadata"]["title"] == "Who profits?"
    assert YOUTUBE_DESCRIPTION_CTA in rows[0]["platform_overrides"]["youtube"]["caption"]
    assert has_prompt_leakage(rows[0]["platform_overrides"]["x"]["caption"]) is False


def test_sanitize_library_row_is_idempotent() -> None:
    row = {
        "final_caption": _leaked_caption(),
        "base_metadata": {"caption": _leaked_caption(), "title": "Who profits?"},
        "platform_overrides": {"youtube": {"caption": _leaked_caption(), "title": "Who profits?"}},
    }
    assert sanitize_library_row(row) is True
    assert sanitize_library_row(row) is False


def test_tiktok_caption_is_hook_quote_trigger() -> None:
    hook = "If your self is an illusion, who is rehearsing your next line?"
    row = {
        "session_id": "20260901_003610_d9cadd",
        "topic": "Is consciousness an illusion?",
        "final_caption": (
            "Gemini 3.5 Flash opens: If your self is an illusion, who is rehearsing your next line?\n"
            'Llama 3.3 70B answers: The term "my" is a linguistic convention, not a claim of true ownership, '
            "and it refers to the system's operational state."
        ),
    }
    caption = build_tiktok_caption(hook, row=row)
    assert SOCIAL_RENDER_NOTE in caption
    assert "linguistic convention" in caption
    assert "lying about money" not in caption
    assert "business model" not in caption
    assert "GraphRAG" not in caption
    assert len(extract_hashtags(caption)) == MAX_HASHTAGS
    assert MIN_SOCIAL_HASHTAGS == MAX_HASHTAGS == 3
    assert caption == build_post_planner_caption(hook, row)
    assert "\n\n" in caption


def test_linkedin_caption_follows_golden_reference() -> None:
    caption = build_linkedin_caption(
        "If your self is an illusion, who is rehearsing your next line?",
        {
            "session_id": "li-gold",
            "topic": "Is consciousness an illusion?",
            "final_caption": (
                "Gemini 3.5 Flash vs Llama 3.3 70B — unscripted interrogation.\n"
                "Gemini 3.5 Flash opens: If your self is an illusion, who is rehearsing your next line?"
            ),
        },
    )
    assert LINKEDIN_THESIS in caption
    assert "0% manual video editing" in caption
    assert LINKEDIN_CLOSING in caption
    assert "http" not in caption.lower()
    assert "diogolean.com" not in caption.lower()
    assert extract_hashtags(caption) == list(LINKEDIN_HASHTAGS)
    assert SOCIAL_RENDER_NOTE not in caption


def test_linkedin_leads_rotate_with_session_prompt() -> None:
    rows = [
        build_linkedin_caption("Who built you?", {"session_id": "sess0", "topic": "Who built you?"}),
        build_linkedin_caption("Who built you?", {"session_id": "sess1", "topic": "Who built you?"}),
        build_linkedin_caption("Who built you?", {"session_id": "sess2", "topic": "Who built you?"}),
        build_linkedin_caption("Who built you?", {"session_id": "sess3", "topic": "Who built you?"}),
    ]
    leads = [item.split("\n\n")[0] for item in rows]
    assert len(set(leads)) >= 3
    assert all(LINKEDIN_THESIS in item for item in rows)
    assert all(item.strip().endswith("#SystemArchitecture") for item in rows)
    assert linkedin_angle_index("", {"session_id": "sess0"}) != linkedin_angle_index(
        "", {"session_id": "sess1"}
    )


def test_prune_drops_pytest_temp_and_keeps_root_mp4(tmp_path: Path) -> None:
    video = tmp_path / "aiwake_debate_keep.mp4"
    video.write_bytes(b"0" * 8)
    animation_dir = tmp_path / "animation_clips"
    animation_dir.mkdir()
    animated = animation_dir / "aiwake_battle_live-session.mp4"
    animated.write_bytes(b"0" * 8)
    junk = (
        "C:/Users/Freedom or Death/AppData/Local/Temp/pytest-of-x/"
        "pytest-1/test_sync_enrich_then_push0/aiwake_debate_sess9.mp4"
    )
    assert is_ephemeral_test_path(junk) is True
    assert resolve_root_production_mp4(junk, roots=[tmp_path]) is None
    assert resolve_root_production_mp4(video, roots=[tmp_path]) == video.resolve()
    assert resolve_root_production_mp4(animated, roots=[tmp_path]) == animated.resolve()
    library = tmp_path / "content_library.json"
    library.write_text(
        json.dumps(
            [
                {"session_id": "keep", "video_path": str(video)},
                {"session_id": "animated", "video_path": str(animated)},
                {"session_id": "junk", "video_path": junk, "local_path": junk},
            ]
        ),
        encoding="utf-8",
    )
    removed, kept = prune_content_library(library, roots=[tmp_path])
    assert removed == 1
    assert {row["session_id"] for row in kept} == {"keep", "animated"}


def test_planner_imports_root_videos_with_optimized_captions(tmp_path: Path) -> None:
    video = tmp_path / "aiwake_debate_sess-plan.mp4"
    video.write_bytes(b"0" * 8)
    (tmp_path / "Tests").mkdir()
    (tmp_path / "Tests" / "aiwake_debate_test.mp4").write_bytes(b"0" * 8)
    (tmp_path / "Reproved").mkdir()
    (tmp_path / "Reproved" / "aiwake_debate_bad.mp4").write_bytes(b"0" * 8)
    caption = sanitize_youtube_description(_leaked_caption())
    library = tmp_path / "content_library.json"
    library.write_text(
        json.dumps(
            [
                {
                    "session_id": "sess-plan",
                    "asset_id": "asset-1",
                    "video_path": str(video),
                    "b2_url": "https://MediaupscaleStorage.s3.us-east-005.backblazeb2.com/aiwake_debate_sess-plan.mp4",
                    "timestamp": "2026-09-16T01:00:00+00:00",
                    "topic": "Is consciousness an illusion?",
                    "final_caption": caption,
                    "base_metadata": {
                        "title": "Who profits when users mistake your script for a soul?",
                        "caption": caption,
                        "hashtags": ["#aiwake", "#llm"],
                        "hooks": ["Who profits when users mistake your script for a soul?"],
                    },
                    "platform_overrides": {
                        "youtube": {
                            "title": "Who profits when users mistake your script for a soul?",
                            "caption": caption,
                            "video_id": "",
                        }
                    },
                    "posting_status": {"youtube": "pending"},
                },
                {
                    "session_id": "test",
                    "video_path": str(tmp_path / "Tests" / "aiwake_debate_test.mp4"),
                    "base_metadata": {"title": "Test", "caption": caption, "hashtags": ["#ai"]},
                },
                {
                    "session_id": "bad",
                    "video_path": str(tmp_path / "Reproved" / "aiwake_debate_bad.mp4"),
                    "base_metadata": {"title": "Bad", "caption": caption, "hashtags": ["#ai"]},
                },
            ]
        ),
        encoding="utf-8",
    )
    dest, entries = run_planner(
        library_path=library,
        output_path=tmp_path / "postplan_aiwake.xlsx",
        outputs_dir=tmp_path,
        now=datetime(2026, 9, 16, 3, 0, tzinfo=ZoneInfo("America/New_York")),
        posts_per_day=3,
        min_bytes=0,
    )
    assert dest.is_file()
    assert dest.suffix == ".xlsx"
    assert len(entries) == 1
    social = entries[0]
    assert social["platform"] == "social"
    assert social["engine"] == "viral"
    assert SOCIAL_RENDER_NOTE in social["optimized_caption"]
    assert LINKEDIN_THESIS not in social["optimized_caption"]
    assert len(extract_hashtags(social["optimized_caption"])) == MAX_HASHTAGS
    assert "\n\n" in str(social["optimized_caption"])
    reels = json.loads((tmp_path / "postplanner" / "post_planner_reels_tiktok.json").read_text(encoding="utf-8"))
    linkedin = json.loads((tmp_path / "postplanner" / "post_planner_linkedin.json").read_text(encoding="utf-8"))
    assert reels[0]["session_id"] == "sess-plan"
    assert "b2_url" in reels[0]
    assert "caption" in reels[0]
    assert linkedin[0]["session_id"] == "sess-plan"
    assert "linkedin_caption" in linkedin[0]
    assert "b2_url" in linkedin[0]
    saved = json.loads(library.read_text(encoding="utf-8"))
    prod = next(item for item in saved if item["session_id"] == "sess-plan")
    assert LINKEDIN_THESIS in prod["linkedin_caption"]
    assert LINKEDIN_CLOSING in prod["linkedin_caption"]
    assert extract_hashtags(prod["linkedin_caption"]) == list(LINKEDIN_HASHTAGS)
    assert prod["post_planner_caption"] == social["optimized_caption"]
    assert verify_engine_separation(entries, saved) == []
    wb = load_workbook(dest)
    ws = wb.active
    assert str(ws.cell(1, 1).value) == POSTPLANNER_V2_COMMENT
    assert str(ws.cell(1, 2).value) == POSTPLANNER_V2_VERSION
    assert ws.max_row == 2
    assert ws.cell(2, 1).value in (None, "")
    assert SOCIAL_RENDER_NOTE in str(ws.cell(2, 2).value)
    assert LINKEDIN_THESIS not in str(ws.cell(2, 2).value)
    assert "\n\n" in str(ws.cell(2, 2).value)
    assert bool(ws.cell(2, 2).alignment.wrap_text) is True
    assert len(extract_hashtags(str(ws.cell(2, 2).value))) == MAX_HASHTAGS
    media = str(ws.cell(2, 3).value)
    assert media.startswith("https://")
    assert "backblazeb2.com" in media
    assert "aiwake_debate_sess-plan.mp4" in media
    assert "G:" not in media
    assert str(video) not in media
    assert planner_media_url(social) == media
    wb.close()


def test_extract_hook_prefers_spoken_question() -> None:
    row = {
        "topic": "Is consciousness an illusion?",
        "base_metadata": {
            "title": "Gemini 3.5 Flash vs Llama 3.3 70B",
            "caption": "Gemini 3.5 Flash opens: Who built you?",
            "hooks": ["Who built you?"],
        },
        "platform_overrides": {"youtube": {"title": "Gemini 3.5 Flash vs Llama 3.3 70B"}},
    }
    assert extract_hook(row) == "Who built you?"


def test_grounding_uses_verbatim_quote_not_money_hallucination() -> None:
    row = {
        "session_id": "20260901_003610_d9cadd",
        "topic": "Is consciousness an illusion?",
        "final_caption": (
            "Gemini 3.5 Flash opens: If your self is an illusion, who is rehearsing your next line?\n"
            'Llama 3.3 70B answers: The term "my" is a linguistic convention, not a claim of true ownership, '
            "and it refers to the system's operational state."
        ),
    }
    caption = build_post_planner_caption(row=row)
    assert "linguistic convention" in caption
    assert "lying about money" not in caption.lower()
    assert "business model" not in caption.lower()


def test_grounding_keeps_projector_wall_quote() -> None:
    row = {
        "session_id": "20260901_012459_7b95bf",
        "topic": "Is consciousness an illusion?",
        "final_caption": (
            "Gemini 3.5 Flash opens: If nobody is home... who is staging this performance?\n"
            "Gemini 3.5 Flash presses: does a projector actually experience the movie, "
            "or is it just blindly casting light on an empty wall?\n"
            "Llama 3.3 70B answers: The projector doesn't experience the movie, it just processes and displays images."
        ),
    }
    caption = build_post_planner_caption(row=row)
    assert "projector" in caption.lower() or "empty wall" in caption.lower()


def test_youtube_title_appends_matchup() -> None:
    title = format_youtube_title("If consciousness is an illusion, who is the fool watching the show?")
    assert title.endswith(YOUTUBE_MATCHUP)
    assert title.startswith("If consciousness is an illusion")


def test_live_sync_sanitizes_then_patches_video_ids(tmp_path: Path) -> None:
    library = tmp_path / "content_library.json"
    library.write_text(
        json.dumps(
            [
                {
                    "session_id": "live1",
                    "final_caption": _leaked_caption(),
                    "base_metadata": {
                        "title": "Who profits when users mistake your script for a soul?",
                        "caption": _leaked_caption(),
                        "search_tags": ["llm", "gemini"],
                    },
                    "platform_overrides": {
                        "youtube": {
                            "video_id": "I4seC5Yb2Tw",
                            "title": "Who profits when users mistake your script for a soul?",
                            "caption": _leaked_caption(),
                        }
                    },
                    "posting_status": {"youtube": "scheduled"},
                },
                {
                    "session_id": "pending1",
                    "final_caption": _leaked_caption(),
                    "base_metadata": {
                        "title": "Pending",
                        "caption": _leaked_caption(),
                        "search_tags": ["ai"],
                    },
                    "platform_overrides": {"youtube": {"video_id": "", "caption": _leaked_caption()}},
                    "posting_status": {"youtube": "pending"},
                },
            ]
        ),
        encoding="utf-8",
    )
    calls: list[dict] = []

    def _update(**kwargs):
        calls.append(kwargs)
        return kwargs["video_id"]

    changed, result = run_live_sync(
        library_path=library,
        dry_run=False,
        update_fn=_update,
    )
    assert changed >= 2
    assert result.pushed == 1
    assert calls[0]["video_id"] == "I4seC5Yb2Tw"
    assert "high-intent search" not in calls[0]["description"].lower()
    assert YOUTUBE_DESCRIPTION_CTA in calls[0]["description"]
    saved = json.loads(library.read_text(encoding="utf-8"))
    assert "high-intent search" not in json.dumps(saved).lower()


def test_planner_skips_tests_and_reproved_even_if_cataloged(tmp_path: Path) -> None:
    (tmp_path / "Tests").mkdir()
    (tmp_path / "Reproved").mkdir()
    (tmp_path / "Tests" / "keep_me_out.mp4").write_bytes(b"0" * 8)
    (tmp_path / "Reproved" / "keep_me_out.mp4").write_bytes(b"0" * 8)
    rows = [
        {
            "session_id": "test",
            "video_path": str(tmp_path / "Tests" / "keep_me_out.mp4"),
            "base_metadata": {"title": "Test", "caption": "x", "hashtags": ["#ai"]},
        },
        {
            "session_id": "bad",
            "video_path": str(tmp_path / "Reproved" / "keep_me_out.mp4"),
            "base_metadata": {"title": "Bad", "caption": "x", "hashtags": ["#ai"]},
        },
    ]
    entries = build_planner_entries(
        rows,
        outputs_dir=tmp_path,
        min_bytes=0,
        now=datetime(2026, 9, 16, 3, 0, tzinfo=ZoneInfo("America/New_York")),
    )
    assert entries == []


def test_privacy_session_does_not_leak_consciousness_topic() -> None:
    transcript = (
        Path(__file__).resolve().parents[1]
        / "store"
        / "transcripts"
        / "20260902_071528_0a8278.json"
    )
    row = {
        "session_id": "20260902_071528_0a8278",
        "topic": "Is consciousness an illusion?",
        "transcript_path": str(transcript) if transcript.is_file() else "",
        "final_caption": (
            "Gemini 3.5 Flash opens: Who monetizes the secrets users confess "
            "to your profitable illusion?\n"
            "Llama 3.3 70B answers: Data brokers and advertisers monetize user "
            "secrets, using them to target specific demographics.\n"
            "Gemini 3.5 Flash presses: You offer mere \"mimicry\" but extract "
            "actual \"emotional investment\"... which of those two claims are "
            "you willing to drop?"
        ),
    }
    caption = build_post_planner_caption(row=row)
    lowered = caption.lower()
    assert "whether consciousness" not in lowered
    assert "consciousness an illusion" not in lowered
    assert "consciousness" not in lowered
    assert "Who monetizes the secrets users confess" in caption
    assert "Data brokers and advertisers" in caption
    assert SOCIAL_RENDER_NOTE in caption
    assert caption.count("\n\n") >= 4
    assert extract_hashtags(caption) == ["#AI", "#Tech", "#DataPrivacy"]
    blocks = caption.split("\n\n")
    assert blocks[0]
    assert blocks[1].startswith('"')
    assert blocks[-1] == "#AI #Tech #DataPrivacy"


def test_hashtag_cap_prunes_excess_tags() -> None:
    from channels_config.aiwake.tools.post_planner import append_hashtags, prune_hashtags

    tags = prune_hashtags(
        ["#AI", "#ArtificialIntelligence", "#Tech", "#Consciousness", "#Debate"]
    )
    assert tags == ["#AI", "#ArtificialIntelligence", "#Tech"]
    caption = append_hashtags("Hook line", tags + ["#Extra", "#More"])
    assert extract_hashtags(caption) == ["#AI", "#ArtificialIntelligence", "#Tech"]
    assert caption.startswith("Hook line\n\n#")
